"""Create the final Chinese DOCX report after formal validation is complete."""

from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DIFFSYNTH_ROOT = Path(r"D:\code\DiffSynth-Studio")
REPORT_DIR = PROJECT_ROOT / "reports"
REPORT_PATH = REPORT_DIR / f"Wan2.1建筑视频LoRA双卡实验报告_{date.today().isoformat()}.docx"
ASSET_DIR = PROJECT_ROOT / "outputs" / "report_assets"
FORMAL_VALIDATION = (
    PROJECT_ROOT / "outputs" / "validation" / "building_train_1912_all_256_cooled"
)
BALANCED_VALIDATION = (
    PROJECT_ROOT / "outputs" / "validation" / "building_train_500_balanced_384_cooled"
)
FORMAL_TRAIN = (
    DIFFSYNTH_ROOT / "models" / "train" / "building_train_1912_all_256_cooled_lora"
)
BALANCED_TRAIN = (
    DIFFSYNTH_ROOT / "models" / "train" / "building_train_500_balanced_384_cooled_lora"
)

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5E6C84"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"
GREEN = "2E7D32"
GOLD = "7A5A00"
RED = "9B1C1C"


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int]):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=110, bottom=110)
            set_cell_borders(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        set_run_font(run, size=9.5, bold=True, color=WHITE)
    repeat_table_header(table.rows[0])
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            if row_index % 2:
                set_cell_shading(cells[index], PALE_BLUE)
            p = cells[index].paragraphs[0]
            if isinstance(value, (int, float)) or len(str(value)) <= 15:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=9.2)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label: str, text: str, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    lead = p.add_run(f"{label}：")
    set_run_font(lead, size=10.5, bold=True, color="000000")
    body = p.add_run(text)
    set_run_font(body, size=10.5, color="000000")


def add_body(doc, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text))
    return p


def add_picture(doc, path: Path, caption: str, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("title", caption)
    picture._inline.docPr.set("descr", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    set_run_font(run, size=9, italic=True, color=MUTED)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, display, end):
        run._r.append(node)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def loss_stats(path: Path):
    values = []
    with path.open("r", encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            if row.get("key") == "loss":
                values.append(float(row["value"]))
    return {
        "rows": len(values),
        "first": values[0],
        "last": values[-1],
        "min": min(values),
        "max": max(values),
        "avg": sum(values) / len(values),
    }


def diagnostics_by_label(path: Path):
    payload = load_json(path)
    return {item["label"]: item["metrics"] for item in payload["results"]}


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    specs = {
        "Heading 1": (16, "000000", 16, 8),
        "Heading 2": (13, "000000", 12, 6),
        "Heading 3": (12, "000000", 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_style = doc.styles["Title"]
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style_p_pr = title_style._element.get_or_add_pPr()
    title_style_border = title_style_p_pr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_p_pr.remove(title_style_border)
    section.header.paragraphs[0].text = ""
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(55)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(kicker.add_run("实验结果报告"), size=11, bold=True, color="000000")
    kicker.paragraph_format.space_after = Pt(14)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    remove_paragraph_border(title)
    set_run_font(title.add_run("Wan2 1 建筑视频 LoRA\n双 RTX 3060 实验报告"), size=27, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(40)
    set_run_font(
        subtitle.add_run("从双卡冒烟验证、500 建筑参数试验到 9560 片段正式训练"),
        size=13,
        color="000000",
    )
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("硬件", "2 × NVIDIA GeForce RTX 3060 12GB"),
            ("基础模型", "Wan2.1-T2V-1.3B / DiffSynth-Studio"),
            ("训练策略", "Windows Gloo DDP、LoRA rank 8、17 帧片段"),
            ("数据划分", "train 1912 / val 239 / test 239；按源 OBJ 分组"),
            ("报告日期", str(date.today())),
        ],
        [2100, 7260],
    )
    add_callout(
        doc,
        "交付范围",
        "本报告记录实际运行结果、失败与恢复过程、检查点质量比较、热安全数据及正式模型建议。",
    )
    doc.add_page_break()


def build_document():
    required = [
        ASSET_DIR / "loss_curves.png",
        ASSET_DIR / "loss_summary.json",
        ASSET_DIR / "gpu_telemetry_summary.json",
        BALANCED_VALIDATION / "video_diagnostics.json",
        BALANCED_VALIDATION / "checkpoint_contact_sheet.png",
        FORMAL_VALIDATION / "video_diagnostics.json",
        FORMAL_VALIDATION / "checkpoint_contact_sheet.png",
        FORMAL_VALIDATION / "checkpoint_selection.json",
        FORMAL_TRAIN / "loss.csv",
        FORMAL_TRAIN / "step-4780.safetensors",
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise SystemExit(f"Formal report inputs are incomplete: {missing}")

    selection = load_json(FORMAL_VALIDATION / "checkpoint_selection.json")
    telemetry = load_json(ASSET_DIR / "gpu_telemetry_summary.json")
    balanced_metrics = diagnostics_by_label(BALANCED_VALIDATION / "video_diagnostics.json")
    formal_metrics = diagnostics_by_label(FORMAL_VALIDATION / "video_diagnostics.json")
    balanced_loss = loss_stats(BALANCED_TRAIN / "loss.csv")
    formal_loss = loss_stats(FORMAL_TRAIN / "loss.csv")
    formal_gpu = {
        int(item["gpu"]): item
        for item in telemetry["groups"]
        if item["stage"] == "formal_train"
    }
    if set(formal_gpu) != {0, 1}:
        raise SystemExit("Formal training telemetry for both GPUs is incomplete.")

    doc = Document()
    configure_styles(doc)
    add_cover(doc)

    doc.add_heading("1 结论摘要", level=1)
    add_callout(
        doc,
        "最终建议",
        f"正式训练推荐使用 {selection['recommended_checkpoint']}。{selection['rationale']}",
        color=GREEN,
    )
    add_body(
        doc,
        "两张 RTX 3060 可以用于同步 LoRA 训练，但不能合并为单张 24GB 显卡。Windows 环境采用 Gloo DDP，正式训练中每张卡保存完整模型副本并同步梯度。",
    )
    add_body(
        doc,
        "500 建筑 384×384 参数试验显示，step-100 的主体居中和局部视角变化最均衡；step-150 之后几何逐渐塌陷，说明仅看 loss 不能判断视觉质量。",
    )
    add_body(
        doc,
        "正式阶段使用冻结的 train 集 1912 个建筑，每个 81 帧视频拆为五个重叠 17 帧窗口，共 9560 个训练片段；val/test 各保留 239 个建筑和 1195 个片段。",
    )

    doc.add_heading("2 数据与实验设计", level=1)
    doc.add_heading("2 1 冻结的数据划分", level=2)
    add_table(
        doc,
        ["集合", "建筑数", "17 帧片段数", "用途"],
        [
            ("train", 1912, 9560, "正式训练"),
            ("validation", 239, 1195, "检查点选择"),
            ("test", 239, 1195, "独立测试"),
            ("总计", 2390, 11950, "不混入训练"),
        ],
        [1500, 1500, 1900, 4460],
    )
    add_body(
        doc,
        "11950 是全库 2390 个视频全部拆成五段后的理论片段总数，不是建筑数量。正式训练只使用 train 的 9560 片段，避免验证集和测试集泄漏。",
    )
    doc.add_heading("2 2 时间片覆盖", level=2)
    for text in (
        "窗口 1：第 1–17 帧",
        "窗口 2：第 17–33 帧",
        "窗口 3：第 33–49 帧",
        "窗口 4：第 49–65 帧",
        "窗口 5：第 65–81 帧",
    ):
        add_bullet(doc, text)

    doc.add_heading("3 实验阶段与运行结果", level=1)
    add_table(
        doc,
        ["阶段", "数据", "分辨率", "优化步", "结果"],
        [
            ("双卡 smoke", "5 片段", "256²", 3, "DDP 与平均 loss 日志通过"),
            ("双卡 smoke", "5 片段", "384²", 3, "显存与速度通过"),
            ("参数试验", "500 balanced", "384²", 250, "step-100 最佳，后期退化"),
            ("正式训练", "1912 / 9560", "256²", 4780, selection["recommended_checkpoint"]),
        ],
        [1800, 1800, 1300, 1300, 3160],
    )
    add_body(
        doc,
        f"500 参数试验：平均 loss {balanced_loss['avg']:.4f}，末步 {balanced_loss['last']:.4f}；正式训练：平均 loss {formal_loss['avg']:.4f}，末步 {formal_loss['last']:.4f}。",
    )
    add_body(
        doc,
        "正式训练的 loss 在早期快速下降，随后围绕较低水平波动并保持稳定；固定参数验证同时表明，loss 更低并不必然意味着检查点的视觉质量更好。",
    )
    doc.add_heading("3 1 检查点判定口径", level=2)
    add_bullet(doc, "所有候选检查点使用相同 prompt、negative prompt、seed 和 30 步推理参数，保证横向比较一致。")
    add_bullet(doc, "主体占比、中心偏移和帧间差异用于辅助量化；最终选择仍以完整 17 帧的连续运动、几何稳定和尺度漂移为主。")
    add_bullet(doc, "推荐检查点不按训练末步自动确定，而是从固定参数验证中选择视觉质量最均衡的中间状态。")
    add_picture(doc, ASSET_DIR / "loss_curves.png", "图 1　各阶段训练 loss 与滑动平均曲线")

    doc.add_heading("4 500 建筑参数试验", level=1)
    add_picture(
        doc,
        BALANCED_VALIDATION / "checkpoint_contact_sheet.png",
        "图 2　500 建筑试验固定 seed 的检查点接触表（每行 5 个时间帧）",
        width=3.25,
    )
    add_table(
        doc,
        ["检查点", "主体占比", "中心偏移", "帧间差异", "判断"],
        [
            ("base", f"{balanced_metrics['base']['foreground_occupancy_mean']:.3f}", f"{balanced_metrics['base']['foreground_centroid_offset_mean']:.3f}", f"{balanced_metrics['base']['mean_frame_difference']:.4f}", "近静态基础形态"),
            ("step-50", f"{balanced_metrics['step-50']['foreground_occupancy_mean']:.3f}", f"{balanced_metrics['step-50']['foreground_centroid_offset_mean']:.3f}", f"{balanced_metrics['step-50']['mean_frame_difference']:.4f}", "主体偏左、运动不足"),
            ("step-100", f"{balanced_metrics['step-100']['foreground_occupancy_mean']:.3f}", f"{balanced_metrics['step-100']['foreground_centroid_offset_mean']:.3f}", f"{balanced_metrics['step-100']['mean_frame_difference']:.4f}", "综合最佳"),
            ("step-150", f"{balanced_metrics['step-150']['foreground_occupancy_mean']:.3f}", f"{balanced_metrics['step-150']['foreground_centroid_offset_mean']:.3f}", f"{balanced_metrics['step-150']['mean_frame_difference']:.4f}", "开始塌陷"),
            ("step-200", f"{balanced_metrics['step-200']['foreground_occupancy_mean']:.3f}", f"{balanced_metrics['step-200']['foreground_centroid_offset_mean']:.3f}", f"{balanced_metrics['step-200']['mean_frame_difference']:.4f}", "几何退化"),
            ("step-250", f"{balanced_metrics['step-250']['foreground_occupancy_mean']:.3f}", f"{balanced_metrics['step-250']['foreground_centroid_offset_mean']:.3f}", f"{balanced_metrics['step-250']['mean_frame_difference']:.4f}", "视角跳变/低细节"),
        ],
        [1400, 1500, 1500, 1500, 3460],
    )

    doc.add_heading("5 正式 9560 片段训练", level=1)
    add_picture(
        doc,
        FORMAL_VALIDATION / "checkpoint_contact_sheet.png",
        "图 3　正式训练固定 prompt、negative prompt、seed 与 30 步推理的检查点接触表",
        width=5.0,
    )
    rows = []
    for label in ("base", "step-500", "step-1000", "step-2000", "step-3000", "step-4000", "step-4780"):
        metrics = formal_metrics[label]
        rows.append(
            (
                label,
                f"{metrics['foreground_occupancy_mean']:.3f}",
                f"{metrics['foreground_centroid_offset_mean']:.3f}",
                f"{metrics['mean_frame_difference']:.4f}",
                "推荐" if label == selection["recommended_checkpoint"] else "对比",
            )
        )
    add_table(
        doc,
        ["检查点", "主体占比", "中心偏移", "帧间差异", "结论"],
        rows,
        [1500, 1500, 1500, 1600, 3260],
    )
    add_callout(doc, "检查点选择", selection["rationale"], color=GREEN)

    doc.add_heading("6 双卡显存与热安全", level=1)
    add_body(
        doc,
        "双卡 smoke 的梯度校验结果为 [80.0, 80.0]，两 rank 同步一致；真实 smoke 也验证了跨 rank 平均 loss 日志不会死锁。",
    )
    add_table(
        doc,
        ["项目", "观察", "处理"],
        [
            ("Windows DDP", "NCCL 不可用", "采用 Gloo + file rendezvous"),
            ("384 训练温度", "未冷却时 GPU1 曾到 92°C", "每步增加 5 秒冷却；后续约 68/80°C"),
            ("页面文件", "双缓存进程触发错误 1455", "改为顺序缓存；配置 D 盘 64–96GB 页面文件"),
            ("磁盘吞吐", "缓存阶段 D 盘常成为瓶颈", "训练阶段使用预计算缓存，双卡持续同步"),
        ],
        [1900, 3000, 4460],
    )
    doc.add_page_break()
    add_table(
        doc,
        ["GPU", "平均利用率", "峰值显存", "最高温度", "最高功耗"],
        [
            (
                f"GPU {gpu}",
                f"{formal_gpu[gpu]['utilization_avg_pct']:.1f}%",
                f"{formal_gpu[gpu]['memory_max_mib'] / 1024:.2f}GB",
                f"{formal_gpu[gpu]['temperature_max_c']:.0f}°C",
                f"{formal_gpu[gpu]['power_max_w']:.1f}W",
            )
            for gpu in (0, 1)
        ],
        [1400, 1800, 1900, 1800, 2460],
    )
    add_body(
        doc,
        "注意：DDP 提升吞吐，但不会把两张 12GB 显卡合并为一张 24GB 显卡。每张卡仍需独立容纳一个训练样本及模型副本。",
    )

    doc.add_heading("7 限制与后续工作", level=1)
    for text in (
        "所有训练 caption 基本同质，模型容易收敛为“平均建筑”，不具备丰富的文本可控差异。",
        "17 帧窗口覆盖完整 81 帧轨迹，但模型学习的是局部连续性，不等同于直接学习 81 帧长程依赖。",
        "当前 LoRA 学到的是建筑领域和旋转视频先验，尚不能单独完成缺失区域几何补全。",
        "下一阶段应加入 incomplete RGB、missing mask、complete RGB 监督与严格相机位姿配对，并设计无 mask、无时序约束等消融实验。",
    ):
        add_bullet(doc, text)

    doc.add_heading("8 可复现实验路径", level=1)
    add_table(
        doc,
        ["产物", "路径"],
        [
            ("正式配置", str(PROJECT_ROOT / "training" / "configs" / "train_1912_all_256_cooled.json")),
            ("正式 LoRA", str(FORMAL_TRAIN)),
            ("正式验证", str(FORMAL_VALIDATION)),
            ("500 LoRA", str(BALANCED_TRAIN)),
            ("500 验证", str(BALANCED_VALIDATION)),
        ],
        [1800, 7560],
    )

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    build_document()
